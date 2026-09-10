from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from openpyxl import Workbook

from data_portal.simplified_workbook import (
    is_simplified_sheet,
    read_simplified_rows,
    validate_simplified_headers,
)


VALID_HEADERS = [
    "escola",
    "codigo_inep",
    "endereco",
    "cep",
    "municipio",
    "estado",
    "tipo",
    "turma",
    "serie",
    "turno",
    "ano_letivo",
    "matricula",
    "nome_aluno",
]


def create_sheet(headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ALUNOS"
    sheet.append(
        headers or VALID_HEADERS
    )

    return sheet


class SimplifiedWorkbookTest(
    SimpleTestCase
):
    def test_detects_simplified_sheet(self):
        sheet = create_sheet()

        self.assertTrue(
            is_simplified_sheet(sheet)
        )

    def test_structured_sheet_is_not_simplified(self):
        sheet = create_sheet([
            "codigo_escola*",
            "codigo_turma*",
            "matricula*",
            "nome_completo*",
        ])

        self.assertFalse(
            is_simplified_sheet(sheet)
        )

    def test_missing_header_is_rejected(self):
        headers = [
            header
            for header in VALID_HEADERS
            if header != "turno"
        ]
        sheet = create_sheet(headers)

        with self.assertRaises(
            ValidationError
        ):
            validate_simplified_headers(
                sheet
            )

    def test_duplicated_header_is_rejected(self):
        headers = list(VALID_HEADERS)
        headers[0] = "codigo_inep"

        sheet = create_sheet(headers)

        with self.assertRaisesMessage(
            ValidationError,
            "cabeçalhos duplicados",
        ):
            validate_simplified_headers(
                sheet
            )

    def test_reads_and_normalizes_row(self):
        sheet = create_sheet()

        sheet.append([
            "  Escola   Municipal Horizonte ",
            26037289,
            " Rua Exemplo, 100  ",
            "56.380-000",
            " Santa Maria da Boa Vista ",
            "Pernambuco",
            "urbano",
            " Única ",
            "2º ano",
            "Matutino",
            2026,
            1542.0,
            "  Ana   Beatriz da Silva ",
        ])

        rows = read_simplified_rows(
            sheet
        )

        self.assertEqual(
            len(rows),
            1,
        )
        self.assertEqual(
            rows[0]["_row"],
            2,
        )
        self.assertEqual(
            rows[0]["escola"],
            "Escola Municipal Horizonte",
        )
        self.assertEqual(
            rows[0]["codigo_inep"],
            "26037289",
        )
        self.assertEqual(
            rows[0]["cep"],
            "56380000",
        )
        self.assertEqual(
            rows[0]["estado"],
            "PE",
        )
        self.assertEqual(
            rows[0]["tipo"],
            "Urbana",
        )
        self.assertEqual(
            rows[0]["turma"],
            "Única",
        )
        self.assertEqual(
            rows[0]["serie"],
            "EF02",
        )
        self.assertEqual(
            rows[0]["turno"],
            "Manhã",
        )
        self.assertEqual(
            rows[0]["ano_letivo"],
            2026,
        )
        self.assertEqual(
            rows[0]["matricula"],
            "1542",
        )
        self.assertEqual(
            rows[0]["nome_aluno"],
            "Ana Beatriz da Silva",
        )

    def test_skips_completely_blank_rows(self):
        sheet = create_sheet()

        sheet.append(
            [None] * len(VALID_HEADERS)
        )
        sheet.append([
            "Escola Horizonte",
            "26037289",
            "",
            "",
            "Município",
            "PE",
            "",
            "A",
            "2",
            "Manhã",
            2026,
            "",
            "Carlos Lima",
        ])

        rows = read_simplified_rows(
            sheet
        )

        self.assertEqual(
            len(rows),
            1,
        )
        self.assertEqual(
            rows[0]["_row"],
            3,
        )

    def test_accepts_headers_with_accents(self):
        headers = list(VALID_HEADERS)
        headers[4] = "Município"
        headers[10] = "Ano letivo"
        headers[12] = "Nome aluno"

        sheet = create_sheet(headers)

        validate_simplified_headers(
            sheet
        )